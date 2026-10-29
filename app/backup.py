"""
backup.py - Keeps your restaurant data safe.

What it does:
  1. AUTOMATIC BACKUP - whenever something changes (new order, event, payment, expense,
     menu change...), a safe copy of restaurant.db is saved within a few minutes.
     It also backs up when the server starts and when it stops.
  2. DAILY COPY      - one copy per day is kept forever in backups/daily/.
  3. EXCEL EXPORT    - every table (orders with totals, events with profit, payments,
     expenses, bookings, menu, staff, discounts...) in ONE Excel file you can open.

Optional: copy every backup to Google Drive too. Install "Google Drive for desktop",
then put this line in your .env file (change the path to your Drive folder):
     BACKUP_COPY_DIR=G:\\My Drive\\Restaurant backups
"""

import json
import os
import shutil
import sqlite3
import threading
import time
import atexit
from datetime import datetime
from pathlib import Path

from app.database import db

PROJECT = Path(__file__).resolve().parent.parent
BACKUP_DIR = PROJECT / "backups"
RECENT_DIR = BACKUP_DIR / "recent"     # the latest copies (old ones are deleted)
DAILY_DIR = BACKUP_DIR / "daily"       # one copy per day, never deleted
EXPORT_DIR = BACKUP_DIR / "excel"      # Excel exports
EXTRA_COPY_DIR = os.getenv("BACKUP_COPY_DIR", "").strip()

KEEP_RECENT = 50          # how many recent copies to keep
WAIT_MINUTES = 5          # back up at most once every 5 minutes after changes

_lock = threading.Lock()
_changed = False
_last_backup = 0.0
_started = False

# Tables or columns that must never go into the Excel file (secrets)
SECRET_TABLES = {"tokens", "sqlite_sequence"}
SECRET_COLUMNS = {"salt", "password_hash", "password"}


# ---------------------------------------------------------------- backup
def database_file() -> Path:
    """Ask SQLite which file it is really using (so we never copy the wrong one)."""
    with db() as conn:
        for row in conn.execute("PRAGMA database_list"):
            if row[1] == "main":
                return Path(row[2])
    raise RuntimeError("Could not find the database file")


def make_backup(reason: str = "manual") -> Path:
    """Save a safe copy of the whole database. Works while the server is running."""
    global _changed, _last_backup
    with _lock:
        RECENT_DIR.mkdir(parents=True, exist_ok=True)
        DAILY_DIR.mkdir(parents=True, exist_ok=True)
        now = datetime.now()
        target = RECENT_DIR / f"restaurant_{now:%Y-%m-%d_%H-%M-%S}_{reason}.db"

        dst = sqlite3.connect(target)
        with db() as conn:
            conn.backup(dst)          # SQLite's own safe copy (handles open files)
        dst.close()

        # Keep today's latest copy in daily/ (one file per day, kept forever)
        shutil.copy2(target, DAILY_DIR / f"restaurant_{now:%Y-%m-%d}.db")

        # Optional second place, e.g. a Google Drive folder
        if EXTRA_COPY_DIR:
            try:
                extra = Path(EXTRA_COPY_DIR)
                extra.mkdir(parents=True, exist_ok=True)
                shutil.copy2(target, extra / f"restaurant_{now:%Y-%m-%d}.db")
            except OSError as err:
                print(f"[backup] Could not copy to {EXTRA_COPY_DIR}: {err}")

        # Delete the oldest recent copies, keep the newest KEEP_RECENT
        recent = sorted(RECENT_DIR.glob("restaurant_*.db"))
        for old in recent[:-KEEP_RECENT]:
            old.unlink(missing_ok=True)

        _changed = False
        _last_backup = time.time()
        print(f"[backup] Saved {target.name}")
        return target


def list_backups() -> dict:
    def info(folder: Path) -> list[dict]:
        return [{"file": p.name, "size_kb": round(p.stat().st_size / 1024, 1),
                 "saved_at": datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds")}
                for p in sorted(folder.glob("*"), reverse=True) if p.is_file()]
    return {"database": str(database_file()),
            "recent": info(RECENT_DIR) if RECENT_DIR.exists() else [],
            "daily": info(DAILY_DIR) if DAILY_DIR.exists() else [],
            "excel": info(EXPORT_DIR) if EXPORT_DIR.exists() else []}


# ---------------------------------------------------------------- automatic
def mark_changed() -> None:
    """Call this after anything changes. The backup runs within WAIT_MINUTES."""
    global _changed
    _changed = True


def _worker() -> None:
    while True:
        time.sleep(30)
        if _changed and time.time() - _last_backup >= WAIT_MINUTES * 60:
            try:
                make_backup("auto")
            except Exception as err:          # never crash the server because of a backup
                print(f"[backup] Automatic backup failed: {err}")


def _on_exit() -> None:
    if _changed:
        try:
            make_backup("shutdown")
        except Exception as err:
            print(f"[backup] Backup on shutdown failed: {err}")


def start_auto_backup() -> None:
    """Start once when the server starts: back up now, then watch for changes."""
    global _started
    if _started:
        return
    _started = True
    try:
        make_backup("startup")
    except Exception as err:
        print(f"[backup] Startup backup failed: {err}")
    threading.Thread(target=_worker, daemon=True, name="auto-backup").start()
    atexit.register(_on_exit)


# ---------------------------------------------------------------- Excel export
def _readable(value):
    """Turn JSON text like {"Biryani": 2} or ["Naan","Dal"] into easy text."""
    if isinstance(value, str) and value[:1] in "[{":
        try:
            value = json.loads(value)
        except ValueError:
            return value
    if isinstance(value, dict):
        return ", ".join(f"{k}: {'; '.join(v) if isinstance(v, list) else v}" for k, v in value.items())
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return value


def _bill(session_id: str, cart: dict) -> dict:
    """Use the app's own bill calculator if it exists; otherwise add up menu prices."""
    try:
        from app.discounts import calculate_bill
        return calculate_bill(session_id, cart)
    except Exception:
        with db() as conn:
            prices = {r["dish"]: r["price"] for r in conn.execute("SELECT dish, price FROM menu")}
        subtotal = sum((prices.get(d) or 0) * q for d, q in cart.items())
        return {"subtotal": subtotal, "discount": 0, "total": subtotal}


def _orders_with_totals(conn) -> list[dict]:
    rows = []
    for r in conn.execute("SELECT * FROM orders ORDER BY updated_at"):
        r = dict(r)
        cart = json.loads(r.get("cart") or "{}")
        bill = _bill(r["session_id"], cart)
        r["items"] = sum(cart.values())
        for k in ("subtotal", "discount", "total"):
            if k in bill:
                r[k] = bill[k]
        rows.append(r)
    return rows


def _events_with_money() -> list[dict]:
    from app.data import get_events
    keep = ("id", "title", "event_type", "status", "date", "time", "customer", "phone",
            "veg_guests", "nonveg_guests", "guests", "veg_package", "nonveg_package",
            "veg_items", "nonveg_items", "total_bill", "investment", "profit", "paid", "balance",
            "description")
    return [{k: e.get(k) for k in keep} for e in get_events()]     # ALL events, past and future


def export_excel() -> Path:
    """Write every table into one Excel file (one sheet each) and return its path."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    path = EXPORT_DIR / f"restaurant_all_data_{datetime.now():%Y-%m-%d_%H-%M}.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="B5471B")

    def add_sheet(title: str, rows: list[dict]):
        ws = wb.create_sheet(title[:31])
        if not rows:
            ws.append(["(no data yet)"])
            return ws
        cols = [c for c in rows[0].keys() if c not in SECRET_COLUMNS]
        ws.append([c.replace("_", " ").title() for c in cols])
        for r in rows:
            ws.append([_readable(r.get(c)) for c in cols])
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = head_font, head_fill
            width = max(len(str(ws.cell(row=j, column=i).value or "")) for j in range(1, min(ws.max_row, 200) + 1))
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    counts = []
    with db() as conn:
        tables = [r[0] for r in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
        nice = {"orders": "Orders", "events": "Events", "event_payments": "Event payments",
                "expenses": "Expenses", "bookings": "Bookings", "menu": "Menu", "staff": "Staff",
                "event_packages": "Buffet packages", "settings": "Settings"}
        first = ["orders", "events", "event_payments", "expenses", "bookings", "menu", "staff"]
        ordered = [t for t in first if t in tables] + [t for t in tables if t not in first]

        for t in ordered:
            if t in SECRET_TABLES:
                continue
            if t == "orders":
                rows = _orders_with_totals(conn)
            elif t == "events":
                rows = _events_with_money()
            else:
                rows = [dict(r) for r in conn.execute(f'SELECT * FROM "{t}"')]
            add_sheet(nice.get(t, t.replace("_", " ").title()), rows)
            counts.append((nice.get(t, t), len(rows)))

        # Money totals for the Summary sheet (all time)
        orders = _orders_with_totals(conn) if "orders" in tables else []
        sales = sum(o.get("total", 0) or 0 for o in orders if o.get("status") == "paid")
        unpaid = sum(o.get("total", 0) or 0 for o in orders if o.get("status") not in ("paid", "cancelled"))
        expenses = conn.execute("SELECT COALESCE(SUM(amount),0) FROM expenses").fetchone()[0] \
            if "expenses" in tables else 0
    events = _events_with_money() if "events" in tables else []
    live = [e for e in events if e.get("status") != "Cancelled"]

    summary.append(["Restaurant data export"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Made on {datetime.now():%d %b %Y, %I:%M %p}"])
    summary.append([])
    summary.append(["All-time money", "Rs."])
    for label, value in [("Restaurant sales collected (paid orders)", sales),
                         ("Orders not paid yet", unpaid),
                         ("Expenses recorded", expenses),
                         ("Event bills (not cancelled)", sum(e.get("total_bill") or 0 for e in live)),
                         ("Event profit", sum(e.get("profit") or 0 for e in live)),
                         ("Event money still to collect", sum(e.get("balance") or 0 for e in live))]:
        summary.append([label, value])
    summary.append([])
    summary.append(["Sheet", "Rows"])
    for name, n in counts:
        summary.append([name, n])
    for row in (4, 12):
        for col in (1, 2):
            summary.cell(row=row, column=col).font = head_font
            summary.cell(row=row, column=col).fill = head_fill
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 16
    for r in summary.iter_rows(min_row=5, max_row=10, min_col=2, max_col=2):
        for c in r:
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")

    wb.save(path)
    return path
